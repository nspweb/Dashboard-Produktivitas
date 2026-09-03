"""Generator laporan 'KOLOM ISIAN' (Hasil Analisis Pre/Post-Test), diporting
dari report_kolom_isian.py asli -- logika inti tidak diubah."""
import re
from datetime import date, datetime
from io import BytesIO
from copy import copy

import openpyxl
from openpyxl.styles import Border, Side, Alignment
from openpyxl.worksheet.page import PageMargins
from openpyxl.utils import get_column_letter

PRIMARY_COLS = [get_column_letter(c) for c in range(3, 19)]      # C..R (16)
SECONDARY_COLS = [get_column_letter(c) for c in range(20, 28)]   # T..AA (8)
DEFAULT_CAPACITY = len(PRIMARY_COLS)
MAX_RESPONDENTS = DEFAULT_CAPACITY + len(SECONDARY_COLS)

MASTER_SHEET = "KOLOM ISIAN"

SECTION_1_ROWS = {"q1": 15, "q2": 16, "q3": 17, "q4": 18, "q5": 19}
SECTION_2_ROWS = {"q6": 25, "q7": 26, "q8": 27, "q9": 28, "q10": 29}
ALL_Q_ROWS = {**SECTION_1_ROWS, **SECTION_2_ROWS}

AVG_ROW_1 = 20
AVG_ROW_2 = 30
HEADER_NUM_ROWS = [13, 23]

TITLE_CELL = "A6"
TRAINING_TITLE_CELL = "A7"
PROGRAM_TITLE_CELL = "A8"
DATE_CELL = "A9"

HARI_ID = {0: "Senin", 1: "Selasa", 2: "Rabu", 3: "Kamis", 4: "Jumat", 5: "Sabtu", 6: "Minggu"}
BULAN_ID = {
    1: "Januari", 2: "Februari", 3: "Maret", 4: "April", 5: "Mei", 6: "Juni",
    7: "Juli", 8: "Agustus", 9: "September", 10: "Oktober", 11: "November", 12: "Desember",
}

TRAINING_TITLE_PREFIX = "PELATIHAN BERBASIS KOMPETENSI KEJURUAN"
PROGRAM_TITLE_PREFIX = "PROGRAM PELATIHAN DASAR"


def format_date_range_id(start: date, end: date) -> str:
    s = f"{HARI_ID[start.weekday()]} {start.day} {BULAN_ID[start.month]}"
    e = f"{HARI_ID[end.weekday()]} {end.day} {BULAN_ID[end.month]} {end.year}"
    return f"TGL {s} S.D {e}"


def format_training_title(extra_name: str) -> str:
    name = re.sub(r"\s+", " ", str(extra_name).strip()).upper()
    if not name:
        return TRAINING_TITLE_PREFIX
    name = re.sub(rf"^{re.escape(TRAINING_TITLE_PREFIX)}\s*", "", name)
    return f"{TRAINING_TITLE_PREFIX} {name}".strip()


def format_program_title(program_name: str) -> str:
    name = re.sub(r"\s+", " ", str(program_name).strip()).upper()
    if not name:
        return PROGRAM_TITLE_PREFIX
    name = re.sub(rf"^{re.escape(PROGRAM_TITLE_PREFIX)}\s*", "", name)
    name = re.sub(r"\s*\([^)]*\)", "", name).strip()
    if "PAKET" not in name:
        m = re.match(r"^(.*\S)\s+(\d+)$", name)
        if m:
            name = f"{m.group(1)} PAKET {m.group(2)}"
    return f"{PROGRAM_TITLE_PREFIX} {name}".strip()


def _parse_date(value):
    if isinstance(value, date):
        return value
    return datetime.strptime(str(value)[:10], "%Y-%m-%d").date()


def _stabilize_master_table(ws):
    thin = Side(style="thin", color="000000")
    for row in range(13, 31):
        for col in range(3, 28):
            cell = ws.cell(row=row, column=col)
            border = cell.border
            cell.border = Border(
                left=border.left if border.left.style else thin,
                right=border.right if border.right.style else thin,
                top=border.top if border.top.style else thin,
                bottom=border.bottom if border.bottom.style else thin,
            )
            cell.alignment = copy(cell.alignment)
            cell.alignment = Alignment(
                horizontal="center", vertical="center",
                wrap_text=cell.alignment.wrap_text,
            )


def _configure_print_layout(ws):
    ws.page_setup.orientation = "landscape"
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_margins = PageMargins(left=0.2, right=0.2, top=0.35, bottom=0.35, header=0.15, footer=0.15)
    ws.print_options.horizontalCentered = True
    ws.print_title_rows = "10:13"


def generate_kolom_isian_report(
    template_path,
    rows: list,
    q_cols: list,
    jenis_test: str,
    training_title_input: str,
    program_name: str,
    tanggal_mulai,
    tanggal_selesai,
) -> tuple:
    n = len(rows)
    if n == 0:
        raise ValueError("Tidak ada responden pada data terfilter untuk laporan ini.")

    tanggal_mulai = _parse_date(tanggal_mulai)
    tanggal_selesai = _parse_date(tanggal_selesai)

    capped = min(n, MAX_RESPONDENTS)
    use_secondary = capped > DEFAULT_CAPACITY
    all_cols = PRIMARY_COLS + (SECONDARY_COLS if use_secondary else [])

    wb = openpyxl.load_workbook(template_path)
    ws = wb[MASTER_SHEET]

    title_label = "HASIL ANALISIS PRE TEST" if jenis_test == "Pre-Test" else "HASIL ANALISIS POST TEST"
    ws[TITLE_CELL] = title_label
    ws[TRAINING_TITLE_CELL] = format_training_title(training_title_input)
    ws[PROGRAM_TITLE_CELL] = format_program_title(program_name)
    ws[DATE_CELL] = format_date_range_id(tanggal_mulai, tanggal_selesai)
    _stabilize_master_table(ws)

    for header_row in HEADER_NUM_ROWS:
        for i, col in enumerate(PRIMARY_COLS):
            ws[f"{col}{header_row}"] = i + 1
        for i, col in enumerate(SECONDARY_COLS):
            ws[f"{col}{header_row}"] = DEFAULT_CAPACITY + i + 1 if use_secondary else None

    data_sorted = sorted(rows, key=lambda r: str(r.get("nama") or ""))[:capped]

    for q, row in ALL_Q_ROWS.items():
        for col in all_cols:
            ws[f"{col}{row}"] = None
        if q not in q_cols:
            continue
        values = [r.get(q) for r in data_sorted]
        for i, col in enumerate(all_cols[:len(values)]):
            v = values[i]
            if v is None:
                continue
            ws[f"{col}{row}"] = int(round(float(v)))

        if use_secondary:
            sum_range = f"C{row}:R{row},T{row}:AA{row}"
        else:
            sum_range = f"C{row}:R{row}"
        ws[f"S{row}"] = f"=IFERROR(SUM({sum_range})/COUNT({sum_range}),0)"

    ws[f"S{AVG_ROW_1}"] = f"=AVERAGE(S{min(SECTION_1_ROWS.values())}:S{max(SECTION_1_ROWS.values())})"
    ws[f"S{AVG_ROW_2}"] = f"=AVERAGE(S{min(SECTION_2_ROWS.values())}:S{max(SECTION_2_ROWS.values())})"

    wb.calculation.fullCalcOnLoad = True

    _configure_print_layout(ws)

    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue(), {
        "jenis_test": jenis_test,
        "respondents_used": len(data_sorted),
        "respondents_total_filtered": n,
        "capped": n > MAX_RESPONDENTS,
        "used_secondary_columns": use_secondary,
    }
