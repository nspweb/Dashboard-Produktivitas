"""Generator laporan resmi (format Kemnaker) untuk sistem REKAP MONEV.
Diporting dari report_generator.py asli -- logika inti TIDAK diubah, hanya
dependensi Streamlit & konversi PDF (LibreOffice) yang dihapus dari file ini."""
import re
import difflib
from pathlib import Path
from io import BytesIO
from copy import copy

import openpyxl
from openpyxl.styles import Border, Side, Font, Alignment
from openpyxl.worksheet.page import PageMargins
from openpyxl.utils import get_column_letter
from openpyxl.drawing.image import Image as ExcelImage

PRIMARY_COLS = [get_column_letter(c) for c in range(3, 19)]   # C..R (16)
SECONDARY_COLS = [get_column_letter(c) for c in range(20, 28)]  # T..AA (8)

DEFAULT_CAPACITY = len(PRIMARY_COLS)                     # 16
MAX_RESPONDENTS = DEFAULT_CAPACITY + len(SECONDARY_COLS)  # 24

MASTER_SHEET = "KOLOM ISIAN"
VIEW_SHEETS = ["Materi Pelatihan", "Program Pelatihan", "INSTRUKTUR", "Penyelenggaraan"]

SHEET_TITLE_ROWS = {
    "KOLOM ISIAN": (6, 29),
    "Histogram": (6, 4),
    "Materi Pelatihan": (6, 29),
    "Program Pelatihan": (6, 29),
    "INSTRUKTUR": (7, 29),
    "Penyelenggaraan": (7, 29),
    "NILAI INST": (6, 6),
}

TITLE_TEXT_CELLS = {
    "Materi Pelatihan": "A6",
    "Program Pelatihan": "A6",
    "INSTRUKTUR": "A7",
    "Penyelenggaraan": "A7",
}

REPORT_TITLE = "HASIL ANALISIS EVALUASI"
HEADER_ROWS = [13, 22, 30, 38]

SECTION_LABELS = {
    "Materi Pelatihan": ["Info mudah didapat", "Pendaftaran mudah", "Petunjuk pendaftaran jelas", "Program jelas"],
    "Program Pelatihan": ["Program menarik", "Program bermanfaat", "Kompetensi meningkat", "Durasi sesuai"],
    "INSTRUKTUR": ["Instruktur menguasai materi", "Kemampuan penyampaian instruktur", "Kemampuan mengelola peserta", "Sikap & teladan instruktur"],
    "Penyelenggaraan": ["Pelayanan petugas", "Jadwal sesuai rencana", "Perlengkapan tepat waktu", "Sarana pelatihan memadai", "Sarana penunjang memadai"],
}

COMMENT_RECAP_HEADERS = [
    "NO.",
    "Komentar dan saran Anda terhadap Pelatihan",
    "Komentar dan saran Anda terhadap Instruktur",
    "Komentar dan saran Anda terhadap penyelenggaraan pelatihan",
    "Keluhan terhadap penyelenggaraan pelatihan",
]


def _normalize(text: str) -> str:
    text = str(text).lower().replace("\n", " ")
    text = re.sub(r"[^a-z0-9]+", " ", text)
    return re.sub(r"\s+", " ", text).strip()


def _find_question_row(ws, question_text: str, cutoff: float = 0.8):
    target = _normalize(question_text)
    best_row, best_ratio = None, 0.0
    for row in ws.iter_rows(min_col=2, max_col=2):
        cell = row[0]
        if not cell.value:
            continue
        norm = _normalize(cell.value)
        if not (norm.startswith("apakah") or norm.startswith("bagaimana")):
            continue
        if norm == target:
            return cell.row
        ratio = difflib.SequenceMatcher(None, norm, target).ratio()
        if norm in target or target in norm:
            ratio = max(ratio, 0.9)
        if ratio > best_ratio:
            best_ratio, best_row = ratio, cell.row
    return best_row if best_ratio >= cutoff else None


def kategori_nilai(score: float) -> str:
    if score >= 5:
        return "SANGAT BAIK"
    if score >= 4:
        return "BAIK"
    if score >= 3:
        return "CUKUP"
    if score >= 2:
        return "KURANG"
    return "TIDAK BAIK"


def generate_comment_recap(
    rows: list,
    info_col: str | None = None,
    comment_col_map: dict | None = None,
    training_comment_label: str = "Komentar/saran program pelatihan",
    instructor_comment_label: str = "Komentar/saran instruktur",
    organizer_comment_label: str = "Komentar/saran penyelenggaraan",
    complaint_label: str = "Keluhan penyelenggaraan",
) -> bytes:
    n = len(rows)
    cmap = comment_col_map or {}

    def get_col(label):
        col = cmap.get(label)
        if col:
            return [("" if r.get(col) is None else str(r.get(col))) for r in rows]
        return [""] * n

    columns_data = [
        list(range(1, n + 1)),
        get_col(training_comment_label),
        get_col(instructor_comment_label),
        get_col(organizer_comment_label),
        get_col(complaint_label),
    ]

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Komentar"

    ws.append(COMMENT_RECAP_HEADERS)
    for cell in ws[1]:
        cell.font = Font(bold=True)
        cell.alignment = Alignment(wrap_text=True, vertical="center", horizontal="center")
    ws.row_dimensions[1].height = 45

    for i in range(n):
        ws.append([columns_data[c][i] for c in range(5)])

    widths = [6.5, 38, 38, 42, 42]
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w

    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical="top")

    ws.page_setup.orientation = "landscape"
    ws.page_setup.paperWidth = "215mm"
    ws.page_setup.paperHeight = "330mm"
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.print_options.horizontalCentered = False
    ws.sheet_view.view = "pageBreakPreview"

    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _format_training_title(training_title: str) -> str:
    name = re.sub(r"\s+", " ", str(training_title).strip()).upper()
    if not name:
        return ""
    name = re.sub(r"^PELATIHAN BERBASIS KOMPETENSI KEJURUAN\s+", "", name)
    name = re.sub(r"^PELATIHAN KOMPETENSI BERBASIS\s+", "", name)
    name = re.sub(r"^PELATIHAN BERBASIS KOMPETENSI\s+", "", name)
    return f"PELATIHAN KOMPETENSI BERBASIS {name}"


def _format_program_title(program_name: str) -> str:
    name = re.sub(r"\s+", " ", str(program_name).strip()).upper()
    if not name:
        return ""
    name = re.sub(r"^PROGRAM PELATIHAN DASAR\s+", "", name)
    name = re.sub(r"\s*\([^)]*\)", "", name)
    name = re.sub(r"\s+", " ", name).strip()
    if "PAKET" not in name:
        match = re.match(r"^(.*\S)\s+(\d+)$", name)
        if match:
            name = f"{match.group(1)} PAKET {match.group(2)}"
    return f"PROGRAM PELATIHAN DASAR {name}"


def _add_title_separator_line(ws, row: int = 6, min_col: int = 1, max_col: int = 29):
    thin_double = Side(style="double", color="000000")
    for col in range(min_col, max_col + 1):
        cell = ws.cell(row=row, column=col)
        existing = cell.border
        cell.border = Border(
            top=thin_double, bottom=existing.bottom,
            left=existing.left, right=existing.right,
        )


def _sync_titles_and_separator_lines(wb):
    for sheet_name, (row, max_col) in SHEET_TITLE_ROWS.items():
        if sheet_name not in wb.sheetnames:
            continue
        _add_title_separator_line(wb[sheet_name], row=row, max_col=max_col)
    for sheet_name, cell_coord in TITLE_TEXT_CELLS.items():
        if sheet_name not in wb.sheetnames:
            continue
        wb[sheet_name][cell_coord] = REPORT_TITLE


def _fix_secondary_header_numbers(ws):
    for row in HEADER_ROWS:
        if ws.cell(row=row, column=19).value != "JUMLAH":
            continue
        for i, col in enumerate(SECONDARY_COLS):
            ws[f"{col}{row}"] = DEFAULT_CAPACITY + i + 1


def _stabilize_master_table(ws):
    thin = Side(style="thin", color="000000")
    for row in range(13, 46):
        for col in range(3, 28):
            cell = ws.cell(row=row, column=col)
            border = cell.border
            cell.border = Border(
                left=border.left if border.left.style else thin,
                right=border.right if border.right.style else thin,
                top=border.top if border.top.style else thin,
                bottom=border.bottom if border.bottom.style else thin,
            )
            if row not in HEADER_ROWS:
                cell.alignment = copy(cell.alignment)
                cell.alignment = Alignment(
                    horizontal="center", vertical="center",
                    wrap_text=cell.alignment.wrap_text,
                )


def _configure_print_layout(ws, orientation: str = "landscape"):
    ws.page_setup.orientation = orientation
    # F4/Folio: 215 x 330 mm. Set physical dimensions because F4 is not
    # represented consistently by Excel/LibreOffice paper-size enums.
    ws.page_setup.paperSize = None
    ws.page_setup.paperWidth = "215mm"
    ws.page_setup.paperHeight = "330mm"
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0 if orientation == "landscape" else 1
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_margins = PageMargins(left=0.2, right=0.2, top=0.35, bottom=0.35, header=0.15, footer=0.15)
    ws.print_options.horizontalCentered = True


def _update_divisor_everywhere(wb, n_respondents: int, use_secondary: bool):
    sum_range = "C{r}:R{r},T{r}:AA{r}" if use_secondary else "C{r}:R{r}"
    for ws in wb.worksheets:
        for row in ws.iter_rows(min_col=19, max_col=19):
            cell = row[0]
            if isinstance(cell.value, str) and cell.value.startswith("=SUM(") and "/12" in cell.value:
                r = cell.row
                cell.value = f"=SUM({sum_range.format(r=r)})/{n_respondents}"


def _style_histogram_chart(wb, training_title: str, section_values=None):
    if "Histogram" not in wb.sheetnames:
        return
    ws = wb["Histogram"]
    if not ws._charts:
        return
    chart = ws._charts[0]
    try:
        runs = chart.title.tx.rich.p[0].r
        label = re.sub(r"\s+", " ", str(training_title).strip()).upper()
        for run in runs:
            if run.t and "NAMA PELATIHAN" in run.t:
                run.t = f" ({label})" if label else " (NAMA PELATIHAN)"
            if run.rPr is not None:
                run.rPr.u = None
                run.rPr.b = True
    except (AttributeError, IndexError):
        pass
    for ser in chart.series:
        ser.graphicalProperties.solidFill = "595959"
        ser.graphicalProperties.line.noFill = True

    if section_values:
        lo = min(section_values)
        hi = max(section_values)
    else:
        lo, hi = 0, 5
    import math
    axis_min = max(0, math.floor(lo * 10) / 10 - 0.2)
    axis_max = min(5, math.ceil(hi * 10) / 10 + 0.2)
    if axis_max - axis_min < 0.5:
        axis_max = min(5, axis_min + 0.5)
    chart.y_axis.scaling.min = round(axis_min, 1)
    chart.y_axis.scaling.max = round(axis_max, 1)
    chart.y_axis.majorUnit = 0.1


def _rewrite_cross_sheet_formulas(wb, row_map: dict, canonical_questions: dict, used_cols: list):
    hide_zero_format = "0;-0;;@"
    for sheet_name in VIEW_SHEETS:
        if sheet_name not in wb.sheetnames:
            continue
        ws = wb[sheet_name]
        allowed_labels = SECTION_LABELS.get(sheet_name, list(row_map.keys()))
        for label in allowed_labels:
            if label not in row_map:
                continue
            master_row = row_map[label]
            question_text = canonical_questions.get(label)
            if not question_text:
                continue
            view_row = _find_question_row(ws, question_text)
            if not view_row:
                continue
            for col in used_cols:
                cell = ws[f"{col}{view_row}"]
                cell.value = f"='{MASTER_SHEET}'!{col}{master_row}"
                cell.number_format = hide_zero_format


def _add_comments_sheet(
    wb,
    template_path: str,
    rows: list,
    comment_col_map: dict | None,
    program_name: str,
    training_title: str,
    date_text: str,
):
    sheet_title = "Komentar dan Saran"
    if sheet_title in wb.sheetnames:
        ws = wb[sheet_title]
    else:
        ws = wb.create_sheet(title=sheet_title)

    # 1. Logo
    logo_path = Path(template_path).with_name("logo_kemnaker.png")
    if logo_path.exists():
        ws._images = []
        logo = ExcelImage(str(logo_path))
        logo.width, logo.height = 105, 92
        logo.anchor = "A1"
        ws.add_image(logo)

    # Helper styles
    font_kop1 = Font(name="Calibri", size=14, bold=True)
    font_kop2 = Font(name="Calibri", size=12, bold=True)
    font_kop3 = Font(name="Calibri", size=16, bold=True)
    font_kop4 = Font(name="Calibri", size=10, bold=False)

    font_title = Font(name="Arial", size=14, bold=True)
    font_header = Font(name="Arial", size=11, bold=True)
    font_data = Font(name="Calibri", size=11, bold=False)

    align_center = Alignment(horizontal="center", vertical="center")
    align_header = Alignment(horizontal="center", vertical="center", wrap_text=True)
    align_no = Alignment(horizontal="center", vertical="top")
    align_comment = Alignment(horizontal="left", vertical="top", wrap_text=True)

    thin_border_side = Side(style="thin", color="000000")
    thin_border = Border(left=thin_border_side, right=thin_border_side, top=thin_border_side, bottom=thin_border_side)
    double_top_side = Side(style="double", color="000000")

    # Row heights
    ws.row_dimensions[1].height = 18.75
    ws.row_dimensions[2].height = 15.75
    ws.row_dimensions[3].height = 21.0
    ws.row_dimensions[4].height = 15.0
    ws.row_dimensions[5].height = 6.0
    ws.row_dimensions[6].height = 20.0
    ws.row_dimensions[7].height = 18.0
    ws.row_dimensions[8].height = 18.0
    ws.row_dimensions[9].height = 18.0
    ws.row_dimensions[10].height = 12.0
    ws.row_dimensions[11].height = 45.0

    # Kop text
    ws["A1"] = "   KEMENTERIAN KETENAGAKERJAAN RI"
    ws["A1"].font = font_kop1
    ws["A1"].alignment = align_center

    ws["A2"] = "            DIREKTORAT JENDERAL PEMBINAAN PELATIHAN DAN PRODUKTIVITAS"
    ws["A2"].font = font_kop2
    ws["A2"].alignment = align_center

    ws["A3"] = "BALAI PELATIHAN VOKASI DAN PRODUKTIVITAS"
    ws["A3"].font = font_kop3
    ws["A3"].alignment = align_center

    ws["A4"] = "Jln. DI Panjaitan No 226 Tlp. 3190427 Kendari 93117"
    ws["A4"].font = font_kop4
    ws["A4"].alignment = align_center

    for r in range(1, 5):
        ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=5)

    # Double separator line on Row 6 top
    for col in range(1, 6):
        cell = ws.cell(row=6, column=col)
        cell.border = Border(top=double_top_side)

    # Title lines
    ws["A6"] = "KOMENTAR DAN SARAN PELATIHAN"
    ws["A6"].font = font_title
    ws["A6"].alignment = align_center

    formatted_training = _format_training_title(training_title) if training_title else "PELATIHAN KOMPETENSI BERBASIS INDUSTRI DAN JASA"
    ws["A7"] = formatted_training
    ws["A7"].font = font_title
    ws["A7"].alignment = align_center

    formatted_program = _format_program_title(program_name) if program_name else "PROGRAM PELATIHAN DASAR"
    ws["A8"] = formatted_program
    ws["A8"].font = font_title
    ws["A8"].alignment = align_center

    ws["A9"] = date_text if date_text else ""
    ws["A9"].font = font_title
    ws["A9"].alignment = align_center

    for r in range(6, 10):
        ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=5)

    # Table Header on Row 11
    headers = [
        "NO.",
        "Komentar dan saran Anda terhadap Pelatihan",
        "Komentar dan saran Anda terhadap Instruktur",
        "Komentar dan saran Anda terhadap penyelenggaraan pelatihan",
        "Keluhan terhadap penyelenggaraan pelatihan",
    ]
    for col_idx, h in enumerate(headers, start=1):
        cell = ws.cell(row=11, column=col_idx, value=h)
        cell.font = font_header
        cell.alignment = align_header
        cell.border = thin_border

    # Resolve column names
    cmap = comment_col_map or {}
    col_pelatihan = cmap.get("Komentar/saran program pelatihan")
    col_instruktur = cmap.get("Komentar/saran instruktur")
    col_penyelenggara = cmap.get("Komentar/saran penyelenggaraan")
    col_keluhan = cmap.get("Keluhan penyelenggaraan")

    if rows and (not col_pelatihan or not col_instruktur or not col_penyelenggara or not col_keluhan):
        row_keys = list(rows[0].keys())
        from .column_match import match_any, COMMENT_CANONICALS
        if not col_pelatihan:
            col_pelatihan = match_any(COMMENT_CANONICALS["Komentar/saran program pelatihan"], row_keys)
        if not col_instruktur:
            col_instruktur = match_any(COMMENT_CANONICALS["Komentar/saran instruktur"], row_keys)
        if not col_penyelenggara:
            col_penyelenggara = match_any(COMMENT_CANONICALS["Komentar/saran penyelenggaraan"], row_keys)
        if not col_keluhan:
            col_keluhan = match_any(COMMENT_CANONICALS["Keluhan penyelenggaraan"], row_keys)

    def _val(row_dict, key):
        if not key or key not in row_dict:
            return ""
        v = row_dict[key]
        if v is None:
            return ""
        s = str(v).strip()
        return "" if s.lower() in ("none", "nan") else s

    for i, r in enumerate(rows):
        row_num = 12 + i
        ws.row_dimensions[row_num].height = None
        row_vals = [
            i + 1,
            _val(r, col_pelatihan),
            _val(r, col_instruktur),
            _val(r, col_penyelenggara),
            _val(r, col_keluhan),
        ]
        for col_idx, val in enumerate(row_vals, start=1):
            cell = ws.cell(row=row_num, column=col_idx, value=val)
            cell.font = font_data
            cell.border = thin_border
            cell.alignment = align_no if col_idx == 1 else align_comment

    widths = [6.5, 38.0, 38.0, 42.0, 42.0]
    for idx, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(idx)].width = w

    ws.page_setup.orientation = "landscape"
    ws.page_setup.paperSize = None
    ws.page_setup.paperWidth = "215mm"
    ws.page_setup.paperHeight = "330mm"
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_margins = PageMargins(left=0.25, right=0.25, top=0.35, bottom=0.35, header=0.15, footer=0.15)
    ws.print_options.horizontalCentered = True
    ws.print_title_rows = "11:11"
    ws.print_area = f"A1:E{11 + len(rows)}"
    ws.sheet_view.view = "pageBreakPreview"


def generate_official_report(
    template_path: str,
    rows: list,
    score_col_map: dict,
    canonical_questions: dict,
    program_name: str,
    training_title: str = "",
    date_text: str = "",
    instructor_name: str = "",
    comment_col_map: dict | None = None,
) -> tuple:
    n = len(rows)
    if n == 0:
        raise ValueError("Tidak ada responden pada data terfilter.")

    capped = min(n, MAX_RESPONDENTS)
    use_secondary = capped > DEFAULT_CAPACITY
    all_cols = PRIMARY_COLS + (SECONDARY_COLS if use_secondary else [])

    wb = openpyxl.load_workbook(template_path)
    ws_master = wb[MASTER_SHEET]

    # EMF logos in the original workbook are dropped by openpyxl on save.
    # Replace the master-sheet header image with a PNG that survives export.
    logo_path = Path(template_path).with_name("logo_kemnaker.png")
    if logo_path.exists():
        ws_master._images = []
        logo = ExcelImage(str(logo_path))
        logo.width, logo.height = 105, 92
        logo.anchor = "B1"
        ws_master.add_image(logo)

    _fix_secondary_header_numbers(ws_master)
    _stabilize_master_table(ws_master)
    _sync_titles_and_separator_lines(wb)
    for sheet in wb.worksheets:
        orient = "portrait" if sheet.title in ["Histogram", "NILAI INST"] else "landscape"
        _configure_print_layout(sheet, orientation=orient)
    ws_master.print_area = [
        "A1:AC19", "A20:AC28", "A29:AC36", "A37:AC45",
    ]

    ws_master["A6"] = REPORT_TITLE
    if training_title:
        ws_master["A7"] = _format_training_title(training_title)
    if program_name:
        ws_master["A8"] = _format_program_title(program_name)
    if date_text:
        ws_master["A9"] = date_text
    if instructor_name:
        ws_master["C31"] = f"Nama Instruktur : {instructor_name}"
        if "INSTRUKTUR" in wb.sheetnames:
            wb["INSTRUKTUR"]["C12"] = f"INSTRUKTUR : {instructor_name}"

    row_map = {}
    for label, question in canonical_questions.items():
        if label not in score_col_map:
            continue
        row = _find_question_row(ws_master, question)
        if row:
            row_map[label] = row

    values_by_label = {}
    for label, row in row_map.items():
        col_name = score_col_map[label]
        raw_values = [r.get(col_name) for r in rows if r.get(col_name) is not None]
        values = [int(round(float(v))) for v in raw_values][:capped]
        values_by_label[label] = values

        for col in all_cols:
            ws_master[f"{col}{row}"] = None
        for i, val in enumerate(values):
            cell = ws_master[f"{all_cols[i]}{row}"]
            cell.value = val
            cell.alignment = cell.alignment.copy(horizontal="center", vertical="center")

    _rewrite_cross_sheet_formulas(wb, row_map, canonical_questions, all_cols)

    section_avgs = []
    for labels in SECTION_LABELS.values():
        means = [sum(values_by_label[l]) / len(values_by_label[l]) for l in labels if l in values_by_label and values_by_label[l]]
        if means:
            section_avgs.append(sum(means) / len(means))
    _style_histogram_chart(wb, training_title or program_name, section_avgs)

    actual_n = max((len(v) for v in values_by_label.values()), default=capped)
    _update_divisor_everywhere(wb, actual_n, use_secondary)

    # Tambahkan lembar Komentar dan Saran sebagai lanjutan setelah NILAI INST
    _add_comments_sheet(
        wb=wb,
        template_path=template_path,
        rows=rows[:actual_n],
        comment_col_map=comment_col_map,
        program_name=program_name,
        training_title=training_title,
        date_text=date_text,
    )

    wb.calculation.fullCalcOnLoad = True

    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue(), {
        "respondents_used": actual_n,
        "respondents_total_filtered": n,
        "capped": n > MAX_RESPONDENTS,
        "used_secondary_columns": use_secondary,
        "questions_matched": len(row_map),
        "questions_total": len(canonical_questions),
        "comments_sheet_added": True,
    }
